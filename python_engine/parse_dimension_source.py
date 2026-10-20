import sys
import os
import re
import json
import time
import csv
from io import StringIO

MONTH_MAP = {
    'januari': 'January', 'februari': 'February', 'maret': 'March',
    'april': 'April', 'mei': 'May', 'juni': 'June',
    'juli': 'July', 'agustus': 'August', 'september': 'September',
    'oktober': 'October', 'november': 'November', 'desember': 'December',
    'pebruari': 'February', 'jan': 'Jan', 'feb': 'Feb', 'mar': 'Mar',
    'apr': 'Apr', 'may': 'May', 'jun': 'Jun', 'jul': 'Jul',
    'aug': 'Aug', 'agu': 'Aug', 'sep': 'Sep', 'oct': 'Oct',
    'okt': 'Oct', 'nov': 'Nov', 'dec': 'Dec', 'des': 'Dec'
}

ROW_ID_KEYWORDS = {
    'no', 'nomor', 'no.', 'no urut', 'bulan', 'month', 'kabupaten/kota', 'kab/kota',
    'kabupaten', 'kota', 'wilayah', 'daerah', 'nama', 'uraian', 'rincian',
    'keterangan', 'kecamatan', 'desa', 'kelurahan', 'stasiun', 'station', 'pos',
    'kelompok umur', 'age group', 'kode pt', 'nama pt', 'deskripsi data', 'deskripsi',
    'instansi', 'unit kerja', 'kantor', 'satuan', 'kode', 'provinsi'
}

KNOWN_EN_TERMS = {
    'male': 'Laki-laki', 'female': 'Perempuan', 'total': 'Jumlah',
    'sex': 'Jenis Kelamin', 'received': 'Masuk', 'finished': 'Diselesaikan',
    'regency': 'Kabupaten', 'city': 'Kota', 'subdistrict': 'Kecamatan',
    'district': 'Kecamatan', 'village': 'Desa', 'month': 'Bulan',
    'year': 'Tahun', 'units': 'Satuan', 'rest': 'Sisa', 'last': 'Akhir',
    'state': 'Negeri', 'private': 'Swasta', 'number': 'Jumlah',
    'pupil': 'Murid', 'student': 'Mahasiswa', 'teacher': 'Guru', 'lecturer': 'Dosen',
    'beginning': 'Awal', 'reported': 'Dilaporkan', 'dismissed': 'Dicabut', 'decided': 'Diputus'
}

def clean_cell_str(val):
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    s = str(val).strip()
    if re.match(r'^-?\d+\.0$', s):
        return s[:-2]
    return s

def extract_cell_lang(val):
    val = (val or "").strip()
    if not val:
        return "", ""
    val_lower = val.lower()
    if val_lower in ['kabupaten/kota', 'kabupaten / kota', 'kab/kota', 'kab / kota']:
        return "Kabupaten/Kota", "Regency/Municipality"

    # 1. Separator baris baru (newline)
    if "\n" in val:
        parts = [p.strip() for p in val.split("\n") if p.strip()]
        if len(parts) >= 2:
            return parts[0], parts[1]
        return parts[0], parts[0]
    # Lindungi teks yang memiliki '/' di dalam tanda kurung seperti (KM2/KMP2) atau kode teknis pendek seperti kd3/d2
    has_slash_in_parens = bool(re.search(r'\([^\)]*\/[^\)]*\)', val))
    is_short_slash_code = bool(re.match(r'^[A-Za-z0-9]{1,4}\s*/\s*[A-Za-z0-9]{1,4}$', val))

    # 2. Separator garis miring (misal: "Angkatan Kerja/ Economically Active" atau "Angkatan Kerja / Economically Active")
    if "/" in val and not re.search(r'\d\s*/\s*\d', val) and not has_slash_in_parens and not is_short_slash_code:
        parts = [p.strip() for p in val.split("/") if p.strip()]
        if len(parts) >= 2:
            if parts[0].lower() == 'lain-lain' and parts[1].lower() == 'lokal':
                return "Lain-lain/Lokal", "Others/Local"
            return parts[0], parts[1]
        return parts[0], parts[0]
    # 3. Cek kamus istilah umum BPS
    val_lower = val.lower()
    if val_lower in KNOWN_EN_TERMS:
        return KNOWN_EN_TERMS[val_lower], val.title()
    for en_k, id_v in KNOWN_EN_TERMS.items():
        if id_v.lower() == val_lower:
            return val, en_k.title()
    return val, val

def read_openpyxl_sheet(ws):
    max_r = ws.max_row or 0
    max_c = ws.max_column or 0
    if max_r == 0 or max_c == 0:
        return []

    grid = [[clean_cell_str(ws.cell(r, c).value) for c in range(1, max_c + 1)] for r in range(1, max_r + 1)]

    # Propagasikan sel merged ke seluruh rentang baris & kolomnya
    try:
        merged_ranges = list(ws.merged_cells.ranges)
        for rng in merged_ranges:
            min_r, min_c = rng.min_row - 1, rng.min_col - 1
            max_r_rng, max_c_rng = min(max_r, rng.max_row), min(max_c, rng.max_col)
            val = grid[min_r][min_c]
            if val != "":
                for r in range(min_r, max_r_rng):
                    for c in range(min_c, max_c_rng):
                        grid[r][c] = val
    except Exception:
        pass

    return grid

def fetch_raw_data(source, sheet_name=None):
    source = source.strip('\r\n\t"')

    # 1. Local Excel file (.xlsx / .xls)
    if os.path.exists(source) and (source.endswith('.xlsx') or source.endswith('.xls')):
        import openpyxl
        wb = openpyxl.load_workbook(source, data_only=True)
        sheet_names = wb.sheetnames

        target_sheet = sheet_names[0]
        if sheet_name and sheet_name in sheet_names:
            target_sheet = sheet_name

        ws = wb[target_sheet]
        raw_rows = read_openpyxl_sheet(ws)
        return {
            "type": "xlsx",
            "sheets": sheet_names,
            "current_sheet": target_sheet,
            "rows": raw_rows
        }

    # 2. Local CSV file
    if os.path.exists(source) and source.endswith('.csv'):
        raw_rows = []
        with open(source, 'r', encoding='utf-8', errors='replace') as f:
            reader = csv.reader(f)
            for row in reader:
                raw_rows.append([clean_cell_str(c) for c in row])
        return {
            "type": "csv",
            "sheets": ["Sheet1"],
            "current_sheet": "Sheet1",
            "rows": raw_rows
        }

    # 3. Google Spreadsheets URL
    if "docs.google.com/spreadsheets" in source:
        import urllib.request
        import ssl
        from io import BytesIO

        match_key = re.search(r"/d/([a-zA-Z0-9-_]+)", source)
        doc_key = match_key.group(1) if match_key else None
        gid = None
        if "gid=" in source:
            gid = source.split("gid=")[1].split("#")[0].split("&")[0].strip()

        now_ts = int(time.time())
        ssl_ctx = ssl._create_unverified_context()

        # Strategi 1: Export format=xlsx
        if doc_key:
            xlsx_url = f"https://docs.google.com/spreadsheets/d/{doc_key}/export?format=xlsx&_t={now_ts}"
            try:
                req = urllib.request.Request(xlsx_url, headers={'User-Agent': 'Mozilla/5.0'})
                with urllib.request.urlopen(req, timeout=12, context=ssl_ctx) as response:
                    raw_bytes = response.read()
                    if len(raw_bytes) > 200 and not raw_bytes.startswith(b"<!DOCTYPE") and not raw_bytes.startswith(b"<html"):
                        import openpyxl
                        wb = openpyxl.load_workbook(BytesIO(raw_bytes), data_only=True)
                        sheet_names = wb.sheetnames
                        target_sheet = sheet_names[0]
                        if sheet_name and sheet_name in sheet_names:
                            target_sheet = sheet_name
                        ws = wb[target_sheet]
                        raw_rows = read_openpyxl_sheet(ws)
                        return {
                            "type": "google_sheets",
                            "sheets": sheet_names,
                            "current_sheet": target_sheet,
                            "rows": raw_rows
                        }
            except Exception:
                pass

        # Strategi 2: Fallback ke CSV export
        if doc_key:
            urls = [
                f"https://docs.google.com/spreadsheets/d/{doc_key}/export?format=csv" + (f"&gid={gid}" if gid else "") + f"&_t={now_ts}",
                f"https://docs.google.com/spreadsheets/d/{doc_key}/gviz/tq?tqx=out:csv" + (f"&gid={gid}" if gid else "") + f"&_t={now_ts}"
            ]
        else:
            csv_link = re.sub(r"/edit.*", "/export?format=csv", source)
            if gid: csv_link += f"&gid={gid}"
            urls = [csv_link + f"&_t={now_ts}"]

        csv_data = None
        for url in urls:
            try:
                req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
                with urllib.request.urlopen(req, timeout=10, context=ssl_ctx) as response:
                    raw = response.read().decode('utf-8', errors='replace')
                    if "<html" not in raw.lower() and "google-site-verification" not in raw.lower():
                        csv_data = raw
                        break
            except Exception:
                continue

        if not csv_data:
            return {"error": "Gagal mengunduh data dari tautan Google Spreadsheet. Pastikan tautan disetel dapat diakses publik (Anyone with the link can view)."}

        reader = csv.reader(StringIO(csv_data))
        raw_rows = [[clean_cell_str(c) for c in row] for row in reader]
        sheet_label = f"Sheet (gid={gid})" if gid else "Sheet Aktif"
        return {
            "type": "google_sheets",
            "sheets": [sheet_label],
            "current_sheet": sheet_label,
            "rows": raw_rows
        }

    return {"error": "Sumber tidak dikenali. Mohon masukkan path file Excel (.xlsx), CSV, atau tautan Google Spreadsheet."}

def is_bracket_col_num(v):
    v = v.strip()
    if not v:
        return False
    # Format penomoran kolom resmi BPS dalam tanda kurung: (1), (2), ( 1 ), (1+2), (1-2), (1)-(2), dll atau -1, -2
    if re.match(r'^(\(\s*[0-9]+(\s*[\+\-\*\/]\s*[0-9]+)*\s*\)|-[0-9]+)$', v):
        return True
    if re.match(r'^\(\s*[0-9]+\s*\)\s*[\+\-\*\/]\s*\(\s*[0-9]+\s*\)$', v):
        return True
    if re.match(r'^\(\s*[a-zA-Z]\s*\)$', v):
        return True
    return False

def is_row_number_str(v):
    v = v.strip()
    if not v:
        return False
    # Angka 0 (jam kerja 0, nilai 0) adalah data baris yang sah, BUKAN nomor urut baris
    if v == "0" or v == "0.0":
        return False
    # Angka dalam tanda kurung seperti (1), (2) adalah penomoran kolom BPS, bukan data baris
    if is_bracket_col_num(v):
        return True
    # Nomor urut baris bertitik murni seperti 1., 2., 01., 02. (tanpa teks kata)
    if re.match(r'^\d+\.$', v):
        return True
    return False

def is_stat_val(v):
    v = v.strip()
    if not v:
        return False
    if is_row_number_str(v):
        return False
    if v in ['-', '–', '—', '...', 'NA', 'N/A', 'nil', 'null']:
        return True
    if re.match(r'^\d+(\s*:\s*\d+(\.\d+)?)?$', v):
        return True
    clean = v.replace(' ', '').replace(',', '').replace('.', '').replace('%', '').replace('(', '').replace(')', '')
    if clean.isdigit() and len(clean) > 0:
        return True
    return False

ROMAN_REGEX = re.compile(r'^(?:M{0,4}(?:CM|CD|D?C{0,3})(?:XC|XL|L?X{0,3})(?:IX|IV|V?I{1,3}))[\.\)]\s+', re.I)
ARABIC_REGEX = re.compile(r'^\d+[\.\)]\s+')
LOWER_REGEX = re.compile(r'^[a-z][\.\)]\s+')
UPPER_REGEX = re.compile(r'^[A-Z][\.\)]\s+')
DASH_REGEX = re.compile(r'^[-–—•]\s+')

def detect_row_hierarchy(raw_rows):
    """
    Mendeteksi level hierarki (0: Bab, 1: Sub Bab, 2: Sub-Sub) untuk rangkaian baris teks.
    Mendukung notasi manual (--/----/spasi) serta pola penomoran resmi BPS:
    - Romawi (I., II.) -> Bab, Angka (1., 2.) -> Sub Bab, Huruf kecil (a., b.) / dash -> Sub-Sub
    - Huruf besar (A., B.) -> Bab, Angka (1., 2.) -> Sub Bab
    - Angka (1., 2.) -> Bab, Huruf kecil (a., b.) / dash -> Sub Bab
    """
    cleaned_rows = [r.strip() for r in raw_rows]
    has_roman = any(ROMAN_REGEX.match(r) for r in cleaned_rows)
    has_arabic = any(ARABIC_REGEX.match(r) for r in cleaned_rows)
    has_upper = any(UPPER_REGEX.match(r) and not ROMAN_REGEX.match(r) for r in cleaned_rows)
    has_lower = any(LOWER_REGEX.match(r) for r in cleaned_rows)
    has_dash = any(DASH_REGEX.match(r) for r in cleaned_rows)

    results = []
    for r in raw_rows:
        val = r.strip()
        level = 0
        clean_v = val

        if r.startswith('----'):
            level = 2
            clean_v = val[4:].strip()
        elif r.startswith('--'):
            level = 1
            clean_v = val[2:].strip()
        elif r.startswith('    '):
            level = 2
            clean_v = val.strip()
        elif r.startswith('  ') or r.startswith('\t'):
            level = 1
            clean_v = val.strip()
        elif has_roman and has_arabic:
            if ROMAN_REGEX.match(val):
                level = 0
            elif ARABIC_REGEX.match(val):
                level = 1
            elif LOWER_REGEX.match(val) or DASH_REGEX.match(val):
                level = 2
            elif any(val.lower().startswith(p) for p in ['jumlah', 'total', 'rata-rata', 'average']):
                level = 0
            else:
                level = 0
        elif has_upper and has_arabic:
            if UPPER_REGEX.match(val):
                level = 0
            elif ARABIC_REGEX.match(val):
                level = 1
            elif LOWER_REGEX.match(val) or DASH_REGEX.match(val):
                level = 2
            elif any(val.lower().startswith(p) for p in ['jumlah', 'total', 'rata-rata', 'average']):
                level = 0
            else:
                level = 0
        elif has_arabic and (has_lower or has_dash):
            if ARABIC_REGEX.match(val):
                level = 0
            elif LOWER_REGEX.match(val) or DASH_REGEX.match(val):
                level = 1
            elif any(val.lower().startswith(p) for p in ['jumlah', 'total', 'rata-rata', 'average']):
                level = 0
            else:
                level = 0

        results.append((clean_v, level))

    return results

def is_col_num_series(row_vals):
    non_empty = [v.strip() for v in row_vals if v.strip()]
    if len(non_empty) < 2:
        return False
    
    # Jika ada kata alfabet >= 4 huruf (seperti nama kabupaten/instansi), bukan nomor kolom!
    for v in non_empty:
        alpha_part = re.sub(r'[^a-zA-Z]', '', v)
        if len(alpha_part) >= 4 and alpha_part.lower() not in ['kolom', 'column']:
            return False

    # Format (1), (2) atau -1, -2
    if all(re.match(r'^(\([0-9]+\)|-[0-9]+|\([0-9]+\s*[\+\-]\s*[0-9]+\))$', v) for v in non_empty):
        return True
    
    # Format angka biasa 1, 2, 3, 4, 5 atau formula
    matched = []
    for v in non_empty:
        m = re.match(r'^([0-9]+)(\s*=\s*[0-9\+\-\*\/\(\)]+)?$', v)
        if m:
            matched.append(int(m.group(1)))
    if len(matched) >= 2 and len(matched) >= len(non_empty) * 0.7:
        if all(x <= 50 for x in matched) and any(matched[i] <= matched[i+1] for i in range(len(matched)-1)):
            return True
    return False

def is_sequence_col(vals):
    cleaned = [v.strip().rstrip('.') for v in vals if v.strip()]
    if len(cleaned) < 2:
        return False
    # PENTING: Jika nilai-nilai ini adalah angka TAHUN (misal 2017, 2018, 2019...), itu adalah dimensi baris waktu, BUKAN nomor urut tabel!
    if sum(1 for v in cleaned if is_year_val(v)) >= len(cleaned) * 0.7:
        return False
    digits = []
    for v in cleaned:
        if v.isdigit():
            digits.append(int(v))
        else:
            return False
    if digits and digits[0] >= 1900 and digits[0] <= 2100:
        return False
    # Deduplikasi nilai berturut-turut yang sama (karena merged cell vertikal pada baris BPS)
    deduped = []
    for d in digits:
        if not deduped or deduped[-1] != d:
            deduped.append(d)
    # Nomor urut baris tabel: baik mulai dari 0, 1, maupun panel lanjutan (misal 16, 21, 31) yang berurutan (+1)
    if len(deduped) >= 2 and all(deduped[i+1] == deduped[i] + 1 for i in range(min(5, len(deduped)-1))):
        return True
    return False

def is_year_val(v):
    return bool(re.match(r'^(19|20)\d{2}(\s*[\-\/]\s*(19|20)?\d{2})?(\s*[\*0-9\(\)a-zA-Z]+)?$', v.strip()))

def detect_structure(grid):
    max_r = len(grid)
    if max_r == 0:
        return 0, 0, 0, [], None
    max_c = max(len(r) for r in grid) if grid else 0
    if max_c == 0:
        return 0, 0, 0, [], None

    for r in range(max_r):
        if len(grid[r]) < max_c:
            grid[r].extend([""] * (max_c - len(grid[r])))

    # 1. Deteksi baris nomor kolom resmi BPS jika ada di 25 baris pertama
    col_num_row_idx = None
    for r in range(min(25, max_r)):
        if is_col_num_series([grid[r][c] for c in range(max_c)]):
            col_num_row_idx = r
            break

    # 2. Cari baris data pertama secara statistik
    data_start = None
    start_search = (col_num_row_idx + 1) if col_num_row_idx is not None else 2
    for r in range(start_search, min(35, max_r)):
        stat_cells = [c for c in range(max_c) if is_stat_val(grid[r][c])]
        
        # Jika hampir seluruh sel numerik adalah angka tahun, ini adalah header tahun!
        year_cells = [c for c in stat_cells if is_year_val(grid[r][c])]
        if len(year_cells) >= 2 and len(year_cells) >= len(stat_cells) * 0.7:
            continue

        has_alpha = any(any(ch.isalpha() for ch in grid[r][c]) for c in range(min(5, max_c)))
        if len(stat_cells) >= 2 and (has_alpha or len(stat_cells) >= 3):
            data_start = r
            break

    if data_start is None:
        data_start = (col_num_row_idx + 1) if col_num_row_idx is not None else min(5, max_r)

    # Jika baris nomor kolom BPS ada dan data_start melompat karena baris teks kategori awal belum memiliki angka statistik,
    # periksa apakah ada baris teks kategori/seksi tepat setelah col_num_row_idx
    if col_num_row_idx is not None and data_start > col_num_row_idx + 1:
        for check_r in range(col_num_row_idx + 1, data_start):
            row_vals = [grid[check_r][c].strip() for c in range(max_c) if grid[check_r][c].strip()]
            if row_vals and any(any(ch.isalpha() for ch in v) for v in row_vals):
                data_start = check_r
                break

    header_end = col_num_row_idx if col_num_row_idx is not None else data_start

    # 3. Cari header_start dengan mundur dari header_end - 1 ke atas
    header_start = header_end - 1 if header_end > 0 else 0
    while header_start > 0:
        prev_r = header_start - 1
        non_empty = [c for c in range(max_c) if grid[prev_r][c].strip()]
        if not non_empty:
            break
        
        row_str = " ".join(grid[prev_r][c].strip() for c in non_empty)
        row_str_lower = row_str.lower()

        is_title = (
            any(row_str_lower.startswith(p) for p in [
                'tabel', 'table', 'lampiran', 'appendix', 'data ', 'statistik',
                'realisasi', 'dinas', 'rekapitulasi', 'posisi ', 'time series',
                'tampilkan data', 'cari nama', 'form '
            ]) or
            'lanjutan tabel' in row_str_lower or 'continued table' in row_str_lower or
            (len(non_empty) <= 3 and len(row_str) > 25 and not any(is_stat_val(grid[prev_r][c]) for c in non_empty)) or
            row_str_lower in ['semua', 'dikdas', 'dikmen']
        )

        if is_title:
            break
        header_start = prev_r

    while header_start < header_end and not any(grid[header_start][c].strip() for c in range(max_c)):
        header_start += 1

    # Cek baris banner seksi di awal data_start (maksimal 2 baris setelah header)
    max_banner_skip = min(data_start + 2, max_r)
    while data_start < max_banner_skip:
        non_empty = [grid[data_start][c].strip() for c in range(max_c) if grid[data_start][c].strip()]
        if not non_empty:
            data_start += 1
            continue
        if len(set(non_empty)) <= 1 and not any(is_stat_val(v) for v in non_empty):
            # Jika baris berikutnya juga memiliki sel teks di kolom yang sama (misal daftar kabupaten di template), jangan lewati!
            if data_start + 1 < max_r and any(grid[data_start + 1][c].strip() for c in range(min(3, max_c))):
                break
            data_start += 1
        else:
            break

    h_rows = [r for r in range(header_start, header_end) if any(grid[r][c].strip() for c in range(max_c))]
    if not h_rows:
        h_rows = [header_start]

    return header_start, header_end, data_start, h_rows, col_num_row_idx

def find_best_row_col(grid, data_start, max_r, max_c, seq_cols, h_rows=None):
    """
    Mencari kolom baris utama pertama dari sebelah kiri tabel (sebelum kolom observasi statistik pertama).
    """
    for c in range(min(10, max_c)):
        if c in seq_cols:
            continue
        sample_vals = []
        for r in range(data_start, min(data_start + 20, max_r)):
            v = grid[r][c].strip()
            if not v:
                continue
            if any(v.lower().startswith(p) for p in ['sumber', 'source', 'catatan', 'note', 'keterangan', 'definisi', 'definition']):
                break
            sample_vals.append(v)
        if not sample_vals:
            continue
        total_vals = len(sample_vals)
        num_start_cnt = sum(1 for v in sample_vals if re.match(r'^\d', v))
        avg_len = sum(len(v) for v in sample_vals) / total_vals
        alpha_cnt = sum(1 for v in sample_vals if any(ch.isalpha() for ch in v) and v.lower() not in ROW_ID_KEYWORDS)
        stat_cnt = sum(1 for v in sample_vals if is_stat_val(v) and not is_year_val(v))
        is_year_series = sum(1 for v in sample_vals if is_year_val(v)) >= total_vals * 0.7

        # Cek apakah header kolom ini mengindikasikan kolom identitas baris / kode / periode
        header_text = ""
        if h_rows:
            header_text = " ".join(grid[r][c].strip().lower() for r in h_rows if r < max_r and grid[r][c].strip())
        is_code_header = any(k in header_text for k in ['kode', 'code', 'kbli', 'klasifikasi', 'golongan', 'komoditas', 'rincian', 'uraian', 'nama', 'item', 'daerah', 'wilayah', 'tahun', 'year', 'bulan', 'month', 'periode', 'period'])

        # Jika kolom ini 100% nomor urut (alpha_cnt == 0, num_start_cnt >= 70%)
        # dan kolom di kanannya memiliki kata alfabet, maka kolom ini nomor urut, bukan nama baris!
        # PENTING: JANGAN anggap nomor urut jika kolom ini adalah DERET TAHUN (2017, 2018..) atau headernya 'tahun'/'year'!
        if not is_year_series and not any(k in header_text for k in ['tahun', 'year']) and alpha_cnt == 0 and num_start_cnt >= total_vals * 0.7 and c + 1 < max_c:
            next_samples = [grid[r][c+1].strip() for r in range(data_start, min(data_start + 20, max_r)) if grid[r][c+1].strip()]
            next_alpha_cnt = sum(1 for v in next_samples if any(ch.isalpha() for ch in v) and v.lower() not in ROW_ID_KEYWORDS)
            if next_alpha_cnt >= 2:
                seq_cols.add(c)
                continue

        # Jika sudah ada seq_cols di sebelah kiri (misal kolom 0 adalah no urut 1, 2, 3), 
        # maka kolom berikutnya yang berisi tahun atau kode BUKAN no urut, melainkan KODE BARIS UTAMA!
        if seq_cols and c > max(seq_cols) and (is_code_header or is_year_series or not is_sequence_col(sample_vals)):
            next_is_stat = False
            if c + 1 < max_c:
                next_samples = [grid[r][c+1].strip() for r in range(data_start, min(data_start + 20, max_r)) if grid[r][c+1].strip()]
                next_stat_cnt = sum(1 for v in next_samples if is_stat_val(v) and not is_year_val(v))
                if len(next_samples) > 0 and next_stat_cnt >= len(next_samples) * 0.6:
                    next_is_stat = True
            if next_is_stat or is_code_header or is_year_series:
                return c

        # Abaikan kolom nomor urut murni (1, 2, 3 atau 01., 02.) HANYA jika bukan kolom kode/tahun dan murni berurutan dari 1
        if num_start_cnt >= total_vals * 0.7 and avg_len < 5 and not is_code_header and not is_year_series:
            cleaned_d = [int(v.split('.')[0]) for v in sample_vals if v.replace('.','').isdigit()]
            if len(cleaned_d) >= 3 and cleaned_d[0] in (0, 1):
                continue

        # Jika sudah bertemu kolom angka data observasi statistik, hentikan pencarian kolom baris pertama dari kiri!
        if stat_cnt >= total_vals * 0.6 and alpha_cnt <= 1 and not is_code_header and not is_year_series:
            break

        unique_texts = set(v.lower() for v in sample_vals if any(ch.isalpha() for ch in v) and v.lower() not in ROW_ID_KEYWORDS)
        if len(unique_texts) >= 2 or alpha_cnt >= 2 or is_code_header or is_year_series:
            return c

    # Fallback pencarian: jika ada seq_cols, kolom tepat setelah seq_cols adalah kandidat terbaik
    if seq_cols:
        candidate = max(seq_cols) + 1
        if candidate < max_c:
            return candidate

    # Fallback pencarian dengan skor jika belum ketemu (hindari memilih kolom observasi statistik panjang)
    best_col = None
    best_score = -1
    for c in range(min(6, max_c)):
        if c in seq_cols:
            continue
        sample_vals = [grid[r][c].strip() for r in range(data_start, min(data_start + 20, max_r)) if grid[r][c].strip()]
        if not sample_vals:
            continue
        total_vals = len(sample_vals)
        alpha_cnt = sum(1 for v in sample_vals if any(ch.isalpha() for ch in v) and v.lower() not in ROW_ID_KEYWORDS)
        unique_texts = set(v.lower() for v in sample_vals if any(ch.isalpha() for ch in v) and v.lower() not in ROW_ID_KEYWORDS)
        stat_cnt = sum(1 for v in sample_vals if is_stat_val(v) and not is_year_val(v))
        avg_len = sum(len(v) for v in sample_vals) / total_vals

        # JANGAN PERNAH memilih kolom yang jelas-jelas data observasi statistik sebagai kolom baris fallback!
        if stat_cnt >= total_vals * 0.8 and alpha_cnt == 0 and avg_len >= 5:
            continue

        score = len(unique_texts) * 50 + alpha_cnt * 10
        if score > best_score:
            best_score = score
            best_col = c

    if best_col is not None:
        return best_col

    return 0

def is_summary_col_header(headers):
    for h in headers:
        if not h: continue
        clean = re.sub(r'[\(\[].*?[\)\]]', '', h).strip().lower()
        clean = clean.split('\n')[0].strip()
        if clean in ['jumlah', 'total', 'jumlah / total', 'total / jumlah', 'sum', 'rata-rata', 'average', 'mean']:
            return True
    return False

def find_continuation_row_cols(grid, primary_col, h_rows, col_num_row_idx, data_start, max_r, max_c):
    if primary_col >= max_c:
        return [primary_col]

    # Cari kolom data observasi pertama untuk mencegah kolom pendamping (sebelum data) dianggap sebagai tabel sambungan samping
    first_data_col = None
    for c in range(primary_col + 1, max_c):
        stat_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if is_stat_val(grid[r][c]) and not is_year_val(grid[r][c]))
        alpha_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if any(ch.isalpha() for ch in grid[r][c]))
        if stat_cnt >= 2 and alpha_cnt <= 1:
            first_data_col = c
            break

    if first_data_col is None:
        for c in range(primary_col + 1, max_c):
            h_vals = [grid[r][c].strip() for r in h_rows if grid[r][c].strip()]
            if any(is_year_val(v) for v in h_vals) or any(v.lower() in MONTH_MAP for v in h_vals):
                first_data_col = c
                break

    if first_data_col is None:
        first_data_col = primary_col + 1

    primary_headers = [grid[r][primary_col].strip().lower() for r in h_rows if r < max_r and grid[r][primary_col].strip()]
    primary_hdr_text = ' '.join(primary_headers)
    primary_col_num = grid[col_num_row_idx][primary_col].strip() if (col_num_row_idx is not None and col_num_row_idx < max_r and primary_col < len(grid[col_num_row_idx])) else ''

    primary_samples = [grid[r][primary_col].strip() for r in range(data_start, min(data_start + 20, max_r)) if grid[r][primary_col].strip()]
    primary_total = len(primary_samples)
    primary_num_cnt = sum(1 for v in primary_samples if v.isdigit() or is_row_number_str(v))
    primary_is_numeric = (primary_num_cnt / primary_total >= 0.7) if primary_total > 0 else False

    cols = [primary_col]

    for c in range(first_data_col + 1, max_c):
        # Harus ada kolom data observasi statistik atau header periode di antara primary_col dan c
        has_sep_stat = any(
            any(is_stat_val(grid[r][mid]) for r in range(data_start, min(data_start + 15, max_r)))
            for mid in range(primary_col + 1, c)
        )
        if not has_sep_stat:
            has_sep_stat = any(
                any(is_year_val(grid[r][mid].strip()) or grid[r][mid].strip().lower() in MONTH_MAP for r in h_rows)
                for mid in range(primary_col + 1, c)
            )
        if not has_sep_stat:
            continue

        c_headers = [grid[r][c].strip().lower() for r in h_rows if r < max_r and grid[r][c].strip()]
        c_hdr_text = ' '.join(c_headers)
        c_col_num = grid[col_num_row_idx][c].strip() if (col_num_row_idx is not None and col_num_row_idx < max_r and c < len(grid[col_num_row_idx])) else ''

        # Jangan anggap kolom nomor urut (1) atau kolom data (3)+ sebagai kolom baris lanjutan
        if re.match(r'^\(?\s*(1|[3-9]|\d{2,})\s*\)?$', c_col_num):
            continue

        # Hindari duplikasi kolom merged horizontal yang sama
        if c > primary_col + 1 and all(grid[r][c].strip() == grid[r][c-1].strip() for r in h_rows):
            prev_cn = grid[col_num_row_idx][c-1].strip() if (col_num_row_idx is not None and col_num_row_idx < max_r and (c-1) < len(grid[col_num_row_idx])) else ''
            if c_col_num == prev_cn:
                continue

        is_same_header = bool(primary_hdr_text and c_hdr_text and (primary_hdr_text == c_hdr_text or primary_hdr_text in c_hdr_text or c_hdr_text in primary_hdr_text))
        is_same_col_num = bool(primary_col_num and is_bracket_col_num(primary_col_num) and c_col_num == primary_col_num)

        if is_same_header or (is_same_col_num and primary_col_num):
            c_samples = [grid[r][c].strip() for r in range(data_start, min(data_start + 20, max_r)) if grid[r][c].strip()]
            c_total = len(c_samples)
            if c_total > 0:
                c_num_cnt = sum(1 for v in c_samples if v.isdigit() or is_row_number_str(v))
                if not primary_is_numeric and (c_num_cnt / c_total >= 0.5):
                    continue
            cols.append(c)

    return cols

def build_dimension_rows(grid, best_row_col_idx, h_rows, col_num_row_idx, data_start, max_r, max_c, shadow_cols=None, en_col=None, clean_headers=True, auto_translate_month=True, auto_hierarchy=True, deduplicate=True):
    if shadow_cols is None:
        shadow_cols = set()
        for c in range(1, max_c):
            is_dup_header = all(grid[r][c].strip() == grid[r][c-1].strip() for r in h_rows)
            is_dup_data = all(grid[r][c].strip() == grid[r][c-1].strip() for r in range(data_start, max_r))
            has_content = any(grid[r][c].strip() for r in range(data_start, max_r))
            if is_dup_header and is_dup_data and has_content:
                shadow_cols.add(c)

    # 1. Deteksi first_data_col
    first_data_col = None
    for c in range(best_row_col_idx + 1, max_c):
        stat_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if is_stat_val(grid[r][c]) and not is_year_val(grid[r][c]))
        alpha_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if any(ch.isalpha() for ch in grid[r][c]))
        if stat_cnt >= 2 and alpha_cnt <= 1:
            first_data_col = c
            break

    if first_data_col is None:
        for c in range(best_row_col_idx + 1, max_c):
            h_vals = [grid[r][c].strip() for r in h_rows if grid[r][c].strip()]
            if any(is_year_val(v) for v in h_vals) or any(v.lower() in MONTH_MAP for v in h_vals):
                first_data_col = c
                break

    if first_data_col is None:
        first_data_col = min(best_row_col_idx + 1, max_c - 1)

    # 2. Deteksi companion_cols (kolom teks rincian/kode antara best_row_col_idx dan first_data_col)
    companion_cols = [c for c in range(best_row_col_idx + 1, first_data_col) if c not in shadow_cols and (en_col is None or c != en_col)]
    comp_col = companion_cols[0] if companion_cols else None

    # Cek apakah comp_col ternyata adalah kolom terjemahan bahasa inggris murni
    if comp_col is not None:
        c_samples = [grid[r][comp_col].strip() for r in range(data_start, min(data_start + 15, max_r)) if grid[r][comp_col].strip()]
        if c_samples:
            match_en_cnt = sum(1 for v in c_samples if v.lower() in KNOWN_EN_TERMS)
            if match_en_cnt >= 2 and match_en_cnt >= len(c_samples) * 0.4:
                comp_col = None

    continuation_row_cols = find_continuation_row_cols(grid, best_row_col_idx, h_rows, col_num_row_idx, data_start, max_r, max_c)
    en_offset = (en_col - best_row_col_idx) if (en_col is not None and en_col >= 0) else None

    raw_candidates = []
    prev_cat = None

    if comp_col is not None and len(continuation_row_cols) == 1 and en_col is None:
        comp_is_unit = any('satuan' in grid[h_r][comp_col].lower() or 'unit' in grid[h_r][comp_col].lower() for h_r in h_rows if comp_col < len(grid[h_r]))
        for r in range(data_start, max_r):
            valA = grid[r][best_row_col_idx].strip()
            valC = grid[r][comp_col].strip() if comp_col < len(grid[r]) else ''
            if not valA and not valC:
                continue
            lowerA = valA.lower()
            if any(lowerA.startswith(p) for p in ['sumber', 'source', 'catatan', 'note', 'keterangan', 'definisi', 'definition']):
                break

            clean_tot = re.sub(r'[^a-z0-9]', '', lowerA)
            is_summary_total = clean_tot in ['jumlah', 'total', 'grandtotal', 'jumlahtotal', 'totaljumlah', 'sum']
            if is_summary_total:
                raw_candidates.append({"raw_str": valA, "en_custom": ""})
                continue

            if comp_is_unit:
                if (valA and not valC) or (valA == valC):
                    if not raw_candidates or raw_candidates[-1]["raw_str"] != valA:
                        raw_candidates.append({"raw_str": valA, "en_custom": ""})
                    prev_cat = valA
                elif valA and valC:
                    clean_unit = valC.strip('()')
                    combined = f"{valA} || {clean_unit}"
                    target_str = ("   " + combined) if prev_cat else combined
                    if not raw_candidates or raw_candidates[-1]["raw_str"] != target_str:
                        raw_candidates.append({"raw_str": target_str, "en_custom": ""})
                elif not valA and valC:
                    clean_unit = valC.strip('()')
                    target_str = ("   " + clean_unit) if prev_cat else clean_unit
                    if not raw_candidates or raw_candidates[-1]["raw_str"] != target_str:
                        raw_candidates.append({"raw_str": target_str, "en_custom": ""})
            else:
                # Non-unit companion (seperti Tabel 12: Komoditi & Provinsi)
                if is_row_number_str(valA):
                    clean_c = valC.strip()
                    combined = f"{valA} || {clean_c}" if clean_c else valA
                    target_str = ("   " + combined) if prev_cat else combined
                    if not raw_candidates or raw_candidates[-1]["raw_str"] != target_str:
                        raw_candidates.append({"raw_str": target_str, "en_custom": ""})
                    continue

                if valA and valC:
                    combined = f"{valA} || {valC}"
                    if not raw_candidates or raw_candidates[-1]["raw_str"] != combined:
                        raw_candidates.append({"raw_str": combined, "en_custom": ""})
                    prev_cat = valA
                elif valA and not valC:
                    if not raw_candidates or raw_candidates[-1]["raw_str"] != valA:
                        raw_candidates.append({"raw_str": valA, "en_custom": ""})
                    prev_cat = valA
                elif not valA and valC:
                    main_val = prev_cat if prev_cat else ""
                    combined = f"{main_val} || {valC}" if main_val else valC
                    if not raw_candidates or raw_candidates[-1]["raw_str"] != combined:
                        raw_candidates.append({"raw_str": combined, "en_custom": ""})
    else:
        for c_idx in continuation_row_cols:
            if c_idx in shadow_cols:
                continue
            target_en_col = (c_idx + en_offset) if (en_offset is not None and c_idx + en_offset < max_c) else None
            for r in range(data_start, max_r):
                if c_idx >= len(grid[r]):
                    continue
                val_ind = grid[r][c_idx].strip()
                if not val_ind or is_bracket_col_num(val_ind):
                    continue
                if clean_headers:
                    if is_col_num_series([val_ind]) or is_row_number_str(val_ind):
                        continue
                    lower_v = val_ind.lower()
                    if lower_v in ROW_ID_KEYWORDS:
                        continue
                    if any(lower_v.startswith(p) for p in ['sumber', 'source', 'catatan', 'note', 'keterangan', 'definisi', 'definition']):
                        break
                val_en_custom = ""
                if target_en_col is not None and target_en_col < len(grid[r]):
                    val_en_custom = grid[r][target_en_col].strip()

                raw_candidates.append({
                    "raw_str": grid[r][c_idx],
                    "en_custom": val_en_custom
                })

    if not raw_candidates:
        return []

    if auto_hierarchy:
        hierarchy_results = detect_row_hierarchy([c["raw_str"] for c in raw_candidates])
    else:
        hierarchy_results = []
        for c in raw_candidates:
            r_str = c["raw_str"]
            level = 0
            clean_text = r_str.strip()
            if r_str.startswith('----'):
                level = 2
                clean_text = r_str[4:].strip()
            elif r_str.startswith('--'):
                level = 1
                clean_text = r_str[2:].strip()
            elif r_str.startswith('    '):
                level = 2
                clean_text = r_str.strip()
            elif r_str.startswith('  ') or r_str.startswith('\t'):
                level = 1
                clean_text = r_str.strip()
            hierarchy_results.append((clean_text, level))

    auto_rows = []
    seen = set()
    current_parent = ""
    for idx, (clean_v, level) in enumerate(hierarchy_results):
        en_custom = raw_candidates[idx]["en_custom"]
        final_ind = clean_v
        final_en = en_custom

        if not final_en:
            if " || " in clean_v:
                p_parts = clean_v.split(" || ")
                ind0, en0 = extract_cell_lang(p_parts[0])
                ind1, en1 = extract_cell_lang(p_parts[1])
                final_ind = f"{ind0} || {ind1}"
                final_en  = f"{en0} || {en1}"
            else:
                split_ind, split_en = extract_cell_lang(clean_v)
                if split_ind and split_en and split_ind != split_en:
                    final_ind = split_ind
                    final_en = split_en
                else:
                    final_en = clean_v

        if auto_translate_month and (not final_en or final_en == final_ind):
            lower_ind = final_ind.lower()
            if lower_ind in MONTH_MAP:
                final_en = MONTH_MAP[lower_ind]

        if level == 0:
            current_parent = final_ind.strip().lower()
            dedup_key = f"0::{final_ind.strip().lower()}"
        else:
            dedup_key = f"{current_parent}::{level}::{final_ind.strip().lower()}"

        if deduplicate:
            if dedup_key in seen:
                continue
            seen.add(dedup_key)

        auto_rows.append({
            "ind": final_ind,
            "en": final_en if final_en else final_ind,
            "level": level
        })

    return auto_rows

def stack_row_continuation_blocks(grid, blocks):
    max_r = len(grid)
    b0_sc, b0_ec = blocks[0]
    b0_grid = [[grid[r][col] for col in range(b0_sc, b0_ec)] for r in range(max_r)]
    h_start, h_end, d_start, h_rows, col_num = detect_structure(b0_grid)

    b0_width = b0_ec - b0_sc

    # Ambil baris header lengkap dari Blok 0
    stacked_grid = [grid[r][b0_sc:b0_ec] for r in range(d_start)]

    # Ambil baris data dari Blok 0
    for r in range(d_start, max_r):
        row_vals = list(grid[r][b0_sc:b0_ec])
        non_empty = [v for v in row_vals if v.strip()]
        if not non_empty:
            continue
        row_txt = ' '.join(non_empty).lower()
        if any(row_txt.startswith(x) for x in ['sumber', 'source', 'catatan', 'note', 'jumlah', 'total', 'lanjutan', 'continued']):
            continue
        stacked_grid.append(row_vals)

    # Ambil baris data dari Blok 1, 2, dst.
    for b_idx, (sc, ec) in enumerate(blocks[1:], 1):
        sub_grid = [[grid[r][col] for col in range(sc, ec)] for r in range(max_r)]
        _, _, b_d_start, _, _ = detect_structure(sub_grid)

        for r in range(b_d_start, max_r):
            row_vals = list(grid[r][sc:ec])
            if len(row_vals) < b0_width:
                row_vals.extend([''] * (b0_width - len(row_vals)))
            elif len(row_vals) > b0_width:
                row_vals = row_vals[:b0_width]

            non_empty = [v for v in row_vals if v.strip()]
            if not non_empty:
                continue
            row_txt = ' '.join(non_empty).lower()
            if any(row_txt.startswith(x) for x in ['sumber', 'source', 'catatan', 'note', 'jumlah', 'total', 'lanjutan', 'continued']):
                continue
            # Lewati jika baris ini nomor kolom resmi BPS seperti (1), (2), (3) atau baris header berulang
            if any(v in ['(1)', '(2)', '(3)', 'kota/city', 'kabupaten/kota', 'provinsi'] for v in [x.lower() for x in non_empty]):
                continue

            stacked_grid.append(row_vals)

    return stacked_grid

def unify_side_by_side_blocks(grid):
    max_r = len(grid)
    if max_r < 3: return grid
    max_c = max(len(r) for r in grid)
    if max_c < 4: return grid

    # 1. Deteksi kolom pemisah kosong (spacer column)
    empty_cols = set()
    for c in range(max_c):
        non_empty = sum(1 for r in range(max_r) if grid[r][c].strip())
        if non_empty == 0:
            empty_cols.add(c)

    if not empty_cols:
        return grid

    # 2. Kelompokkan kolom menjadi kandidat blok tabel horizontal
    blocks = []
    curr_start = 0
    c = 0
    while c < max_c:
        if c in empty_cols:
            if c > curr_start:
                blocks.append((curr_start, c))
            while c < max_c and c in empty_cols:
                c += 1
            curr_start = c
        else:
            c += 1
    if curr_start < max_c:
        blocks.append((curr_start, max_c))

    if len(blocks) < 2:
        return grid

    # 3. Deteksi apakah blok-blok samping merupakan Lanjutan Baris (Row Continuation / Continued Table BPS)
    is_row_continuation = False

    # Indikator A: Teks "Lanjutan" / "Continued" / "Bersambung" di baris header awal blok 1 dst
    for b_idx, (sc, ec) in enumerate(blocks[1:], 1):
        for r in range(min(6, max_r)):
            txt = ' '.join(grid[r][col] for col in range(sc, ec)).lower()
            if any(k in txt for k in ['lanjutan', 'continued', 'bersambung']):
                is_row_continuation = True
                break
        if is_row_continuation:
            break

    # Indikator B: Penomoran urut baris di awal blok bersambung (misal blok 0 dimulai dari 1, blok 1 dari 31, dst)
    if not is_row_continuation:
        first_row_nums = []
        for b_idx, (sc, ec) in enumerate(blocks):
            found_num = None
            for r in range(min(12, max_r)):
                val = grid[r][sc].strip()
                if val.isdigit() and int(val) > 0:
                    found_num = int(val)
                    break
            first_row_nums.append(found_num)

        if len(first_row_nums) >= 2 and first_row_nums[0] is not None and first_row_nums[1] is not None:
            if first_row_nums[1] > first_row_nums[0]:
                is_row_continuation = True

    if is_row_continuation:
        return stack_row_continuation_blocks(grid, blocks)

    parsed_blocks = []
    for b_idx, (start_c, end_c) in enumerate(blocks):
        sub_grid = [[grid[r][col] for col in range(start_c, end_c)] for r in range(max_r)]
        h_start, h_end, d_start, h_rows, col_num = detect_structure(sub_grid)
        b_max_c = end_c - start_c

        # Bersihkan baris judul (tabel/table/dll) dari h_rows agar tidak merusak nama kolom
        clean_h_rows = []
        for r in h_rows:
            row_text = ' '.join(sub_grid[r]).lower()
            if any(row_text.startswith(p) for p in ['tabel', 'table', 'volume dan nilai', 'volume and value', 'lampiran', 'appendix']):
                continue
            clean_h_rows.append(r)
        if not clean_h_rows:
            clean_h_rows = h_rows

        # Deteksi tahun/periode di header blok
        period = None
        for r in clean_h_rows:
            for col in range(b_max_c):
                v = sub_grid[r][col].strip()
                if is_year_val(v):
                    period = v
                    break
            if period: break

        # Eliminasi kolom bayangan (shadow columns akibat merged cell horizontal)
        cols_to_remove = set()
        for col in range(b_max_c - 1):
            if all(sub_grid[r][col].strip() == sub_grid[r][col+1].strip() for r in clean_h_rows):
                # PENTING: Jangan hapus jika data rows di col dan col+1 memiliki nilai yang berbeda secara independen!
                # (misal col=Tahun dan col+1=Bulan yang di headernya di-merge Tahun/Year)
                diff_count = sum(1 for r in range(d_start, min(d_start + 25, max_r)) 
                                 if sub_grid[r][col].strip() and sub_grid[r][col+1].strip() 
                                 and sub_grid[r][col].strip() != sub_grid[r][col+1].strip())
                if diff_count >= 2:
                    continue
                cnt_c = sum(1 for r in range(d_start, len(grid)) if sub_grid[r][col].strip())
                cnt_next = sum(1 for r in range(d_start, len(grid)) if sub_grid[r][col+1].strip())
                if cnt_c < cnt_next:
                    cols_to_remove.add(col)
                else:
                    cols_to_remove.add(col+1)
        compact_cols = [col for col in range(b_max_c) if col not in cols_to_remove]

        # Pisahkan kolom baris (dimensi kategori) vs kolom data (observasi statistik)
        row_cols = []
        data_cols = []
        for col in compact_cols:
            cn = sub_grid[col_num][col].strip() if (col_num is not None and col_num < max_r) else ''
            if cn in ('(1)', '(2)', '1', '2'):
                row_cols.append(col)
            elif re.match(r'^\(?\s*([3-9]|\d{2,})\s*\)?$', cn):
                data_cols.append(col)
            else:
                stat_cnt = sum(1 for r in range(d_start, min(d_start + 15, max_r)) if is_stat_val(sub_grid[r][col]) and not is_year_val(sub_grid[r][col]))
                alpha_cnt = sum(1 for r in range(d_start, min(d_start + 15, max_r)) if any(ch.isalpha() for ch in sub_grid[r][col]))
                if stat_cnt >= 2 and alpha_cnt <= 1:
                    data_cols.append(col)
                else:
                    row_cols.append(col)

        if not data_cols or not row_cols:
            return grid

        col_headers = []
        for dc in data_cols:
            hdr_parts = [sub_grid[r][dc].strip() for r in clean_h_rows if sub_grid[r][dc].strip() and not is_year_val(sub_grid[r][dc].strip())]
            hdr = ' / '.join(hdr_parts) if hdr_parts else f'Data {dc}'
            col_headers.append(hdr)

        row_headers = []
        for rc in row_cols:
            hdr_parts = [sub_grid[r][rc].strip() for r in clean_h_rows if sub_grid[r][rc].strip() and not is_year_val(sub_grid[r][rc].strip())]
            hdr = hdr_parts[-1] if hdr_parts else f'Baris {rc}'
            row_headers.append(hdr)

        b_data = {}
        curr_row_keys = [''] * len(row_cols)
        for r in range(d_start, max_r):
            for i, rc in enumerate(row_cols):
                val = sub_grid[r][rc].strip()
                if not val and rc > 0:
                    val = sub_grid[r][rc-1].strip()
                if val:
                    curr_row_keys[i] = val
            key = tuple(curr_row_keys)
            if not any(key): continue
            if any(any(k.lower().startswith(x) for x in ['sumber', 'source', 'catatan', 'note', 'jumlah', 'total']) for k in key):
                continue
            vals = [sub_grid[r][dc].strip() for dc in data_cols]
            b_data[key] = vals

        parsed_blocks.append({
            'period': period,
            'clean_h_rows': clean_h_rows,
            'sub_grid': sub_grid,
            'row_headers': row_headers,
            'col_headers': col_headers,
            'row_cols': row_cols,
            'data_cols': data_cols,
            'data': b_data
        })

    num_row_dims = len(parsed_blocks[0]['row_headers'])
    if any(len(pb['row_headers']) != num_row_dims for pb in parsed_blocks):
        return grid

    base_row_headers = parsed_blocks[0]['row_headers']

    all_keys = []
    for pb in parsed_blocks:
        for k in pb['data'].keys():
            if k not in all_keys:
                all_keys.append(k)

    if not all_keys:
        return grid

    # Cek apakah blok-blok memiliki periode waktu yang nyata dan berbeda
    distinct_periods = [pb['period'] for pb in parsed_blocks if pb['period']]
    has_distinct_periods = len(distinct_periods) == len(parsed_blocks) and len(set(distinct_periods)) > 1

    if not has_distinct_periods:
        # Jika bukan beda periode (hanya tabel lebar dengan indikator berbeda berdampingan):
        # Rekonstruksi baris header asli tanpa menyisipkan fake "Periode 1", "Periode 2"
        unified_grid = []
        h_rows_all = parsed_blocks[0]['clean_h_rows']
        for hr in h_rows_all:
            h_row = [parsed_blocks[0]['sub_grid'][hr][rc] for rc in parsed_blocks[0]['row_cols']]
            for pb in parsed_blocks:
                h_row.extend([pb['sub_grid'][hr][dc] for dc in pb['data_cols']])
            unified_grid.append(h_row)

        for k in all_keys:
            row_vals = list(k)
            for pb in parsed_blocks:
                d_vals = pb['data'].get(k, [''] * len(pb['data_cols']))
                row_vals.extend(d_vals)
            unified_grid.append(row_vals)

        return unified_grid

    hdr_0 = list(base_row_headers)
    for pb in parsed_blocks:
        hdr_0.extend([pb['period'] or f'Periode {b_idx+1}'] * len(pb['col_headers']))

    hdr_1 = list(base_row_headers)
    for pb in parsed_blocks:
        hdr_1.extend(pb['col_headers'])

    hdr_2 = [f'({i+1})' for i in range(num_row_dims)]
    cur_n = num_row_dims + 1
    for pb in parsed_blocks:
        for _ in pb['col_headers']:
            hdr_2.append(f'({cur_n})')
            cur_n += 1

    unified_grid = [hdr_0, hdr_1, hdr_2]
    for k in all_keys:
        row_vals = list(k)
        for pb in parsed_blocks:
            d_vals = pb['data'].get(k, [''] * len(pb['col_headers']))
            row_vals.extend(d_vals)
        unified_grid.append(row_vals)

    return unified_grid

def inspect_source(source, sheet_name=None, row_col_override=None):
    res = fetch_raw_data(source, sheet_name)
    if "error" in res:
        return res

    grid = res.get("rows", [])
    if not grid:
        return {"error": "Data kosong (tidak ditemukan baris dalam sheet)."}

    # Unifikasi otomatis tabel multi-block (tabel berdampingan horizontal beda periode/tahun)
    grid = unify_side_by_side_blocks(grid)

    max_r = len(grid)
    max_c = max(len(r) for r in grid) if grid else 0
    if max_c == 0:
        return {"error": "Sheet tidak memiliki kolom."}

    # Normalisasi panjang setiap baris agar seragam
    for r in range(max_r):
        if len(grid[r]) < max_c:
            grid[r].extend([""] * (max_c - len(grid[r])))

    header_start, header_end, data_start, h_rows, col_num_row_idx = detect_structure(grid)

    # Deteksi kolom nomor urut baris (1, 2, 3...)
    seq_cols = set()
    for c in range(min(4, max_c)):
        sample_vals = [grid[r][c] for r in range(data_start, min(data_start + 10, max_r))]
        if is_sequence_col(sample_vals):
            seq_cols.add(c)
            break

    # Cari kolom teks nama baris utama (atau gunakan override dari input user jika diberikan)
    if row_col_override is not None and str(row_col_override).strip() not in ("", "None", "-1"):
        best_row_col_idx = int(row_col_override)
    else:
        best_row_col_idx = find_best_row_col(grid, data_start, max_r, max_c, seq_cols, h_rows)

    # Cari kolom data statistik pertama (kolom numerik observasi setelah best_row_col_idx)
    first_stat_col = None
    # 1. Dari baris nomor kolom resmi BPS jika ada (nomor (3) ke atas adalah kolom data)
    if col_num_row_idx is not None and col_num_row_idx < max_r:
        for c in range(best_row_col_idx + 1, max_c):
            cn = grid[col_num_row_idx][c].strip()
            if re.match(r'^\(?\s*([3-9]|\d{2,})\s*\)?$', cn):
                first_stat_col = c
                break

    # 2. Dari data statistik observasi
    if first_stat_col is None:
        for c in range(best_row_col_idx + 1, max_c):
            sample_vals = [grid[r][c] for r in range(data_start, min(data_start + 10, max_r))]
            if is_sequence_col(sample_vals):
                continue
            stat_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if is_stat_val(grid[r][c]) and not is_year_val(grid[r][c]))
            alpha_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if any(ch.isalpha() for ch in grid[r][c]))
            if stat_cnt >= 2 and alpha_cnt <= 1:
                first_stat_col = c
                break

    if first_stat_col is None:
        for c in range(best_row_col_idx + 1, max_c):
            h_vals = [grid[r][c].strip() for r in h_rows if grid[r][c].strip()]
            if any(is_year_val(v) for v in h_vals) or any(v.lower() in MONTH_MAP for v in h_vals):
                first_stat_col = c
                break

    if first_stat_col is None:
        first_stat_col = best_row_col_idx + 1

    # Kolom identitas baris: best_row_col_idx dan seluruh kolom teks/atribut sebelum first_stat_col
    row_identifier_cols = {best_row_col_idx}
    for c in range(first_stat_col):
        cn = grid[col_num_row_idx][c].strip() if (col_num_row_idx is not None and col_num_row_idx < max_r and c < len(grid[col_num_row_idx])) else ''
        if re.match(r'^\(?\s*([3-9]|\d{2,})\s*\)?$', cn):
            continue
        alpha_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if any(ch.isalpha() for ch in grid[r][c]))
        sample_vals = [grid[r][c] for r in range(data_start, min(data_start + 10, max_r))]
        if alpha_cnt >= 2 or is_sequence_col(sample_vals) or any(is_row_number_str(v) for v in sample_vals if v) or cn in ('(1)', '(2)', '1', '2'):
            row_identifier_cols.add(c)

    # Jika kolom di sebelah kanan best_row_col memiliki header sama persis dengan best_row_col (merged cell di baris header dan bukan kolom angka)
    for c in range(best_row_col_idx + 1, max_c):
        same_h = any(grid[r][c].strip() == grid[r][best_row_col_idx].strip() and grid[r][best_row_col_idx].strip() != "" for r in h_rows)
        stat_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if is_stat_val(grid[r][c]) and not is_year_val(grid[r][c]))
        alpha_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if any(ch.isalpha() for ch in grid[r][c]))
        if same_h and stat_cnt < 2 and alpha_cnt >= 2:
            row_identifier_cols.add(c)
        else:
            break

    # Deteksi tabel bersambung samping (side-by-side continuation)
    continuation_cols = find_continuation_row_cols(grid, best_row_col_idx, h_rows, col_num_row_idx, data_start, max_r, max_c)
    for c in continuation_cols:
        row_identifier_cols.add(c)
        if c > 0 and (c - 1) not in row_identifier_cols:
            c_samples = [grid[r][c].strip() for r in range(data_start, min(data_start + 6, max_r))]
            c_non_empty = [v for v in c_samples if v]
            # Hanya periksa kolom sebelumnya jika kolom c saat ini BUKAN kolom nomor urut itu sendiri
            if not is_sequence_col(c_non_empty):
                prev_samples = [grid[r][c-1].strip() for r in range(data_start, min(data_start + 6, max_r))]
                prev_non_empty = [v for v in prev_samples if v]
                prev_headers = [grid[r][c-1].strip() for r in h_rows if grid[r][c-1].strip()]
                prev_cn = grid[col_num_row_idx][c-1].strip() if (col_num_row_idx is not None and col_num_row_idx < max_r and (c-1) < len(grid[col_num_row_idx])) else ''
                if len(prev_non_empty) >= 2 and not is_summary_col_header(prev_headers):
                    if is_sequence_col(prev_non_empty) or prev_cn in ('(1)', '1'):
                        row_identifier_cols.add(c - 1)

    # Identifikasi kolom baris resmi BPS berdasarkan nomor kolom (1) dan (2)
    if col_num_row_idx is not None and col_num_row_idx < max_r:
        for c in range(max_c):
            cn = grid[col_num_row_idx][c].strip() if c < len(grid[col_num_row_idx]) else ''
            if cn in ('(1)', '(2)', '1', '2'):
                h_vals = [grid[r][c].strip() for r in h_rows if grid[r][c].strip()]
                is_period_or_year = any(is_year_val(v) for v in h_vals)
                stat_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if is_stat_val(grid[r][c]) and not is_year_val(grid[r][c]))
                alpha_cnt = sum(1 for r in range(data_start, min(data_start + 15, max_r)) if any(ch.isalpha() for ch in grid[r][c]))
                if (is_period_or_year or (stat_cnt >= 2 and alpha_cnt <= 1)) and (c not in seq_cols):
                    row_identifier_cols.discard(c)
                else:
                    row_identifier_cols.add(c)
            elif re.match(r'^\(?\s*([3-9]|\d{2,})\s*\)?$', cn):
                row_identifier_cols.discard(c)

    # Pastikan kolom ringkasan/total data (Jumlah/Total) TIDAK PERNAH masuk sebagai row_identifier_cols
    for c in range(max_c):
        col_headers = [grid[r][c].strip() for r in h_rows if grid[r][c].strip()]
        if is_summary_col_header(col_headers):
            row_identifier_cols.discard(c)

    # 2.5 Deteksi Kolom Bayangan (Shadow Columns) akibat merged cell horizontal di Excel
    shadow_cols = set()
    for c in range(1, max_c):
        is_dup_header = all(grid[r][c].strip() == grid[r][c-1].strip() for r in h_rows)
        is_dup_data = all(grid[r][c].strip() == grid[r][c-1].strip() for r in range(data_start, max_r))
        has_content = any(grid[r][c].strip() for r in range(data_start, max_r))
        if is_dup_header and is_dup_data and has_content:
            shadow_cols.add(c)

    row_identifier_cols.difference_update(shadow_cols)

    # 3. Deteksi Baris Periode Waktu di Header
    period_row_indices = set()
    year_row_idx = None
    month_row_idx = None

    data_cols = [c for c in range(max_c) if c not in row_identifier_cols and c not in shadow_cols]
    for r in h_rows:
        row_vals = [grid[r][c].strip() for c in data_cols if grid[r][c].strip()]
        if not row_vals: continue
        total = len(row_vals)

        # Cek Tahun
        y_count = sum(1 for v in row_vals if re.match(r'^(19|20)\d{2}(\s*[\-\/]\s*(19|20)?\d{2})?$', v))
        if y_count >= 2 and (y_count >= total * 0.35 or y_count >= 2):
            period_row_indices.add(r)
            year_row_idx = r

        # Cek Bulan
        m_count = sum(1 for v in row_vals if v.split('\n')[0].strip().lower() in MONTH_MAP or v.split('/')[0].strip().lower() in MONTH_MAP)
        if m_count >= 2 and (m_count >= total * 0.35 or m_count >= 2):
            period_row_indices.add(r)
            month_row_idx = r

    # 4. Baris Header Khusus Karakteristik (Tanpa baris periode waktu dan tanpa baris nomor kolom BPS)
    col_h_rows = [r for r in h_rows if r not in period_row_indices and r != col_num_row_idx and any(grid[r][c].strip() for c in data_cols)]
    if not col_h_rows and h_rows:
        col_h_rows = [r for r in h_rows if r != col_num_row_idx]

    # 5. Susun Columns Info & Hierarki Cerdas (Mendukung hingga 3 Tingkatan Kolom Asimetris)
    # Deteksi apakah ada baris translasi bahasa murni (misal Row r = Indo, Row r+1 = English)
    en_translation_rows = set()
    if len(col_h_rows) >= 2:
        for i in range(len(col_h_rows) - 1):
            r_top = col_h_rows[i]
            r_bot = col_h_rows[i+1]
            match_en_count = 0
            valid_col_count = 0
            has_span = False
            for c in data_cols:
                v_top = grid[r_top][c].strip()
                v_bot = grid[r_bot][c].strip()
                if not v_top or not v_bot:
                    continue
                valid_col_count += 1
                if c + 1 < max_c and grid[r_top][c+1].strip() == v_top and grid[r_bot][c+1].strip() != v_bot:
                    has_span = True
                if v_bot.lower() in KNOWN_EN_TERMS or (v_top.lower() in KNOWN_EN_TERMS.values()):
                    match_en_count += 1
            if not has_span and valid_col_count > 0 and match_en_count >= valid_col_count * 0.5:
                en_translation_rows.add(r_bot)

    columns_info = []
    distinct_parents = []
    seen_parents = set()

    for c in range(max_c):
        if c in shadow_cols:
            continue
        col_letter = chr(65 + c) if c < 26 else f"Col{c+1}"
        is_row_id = (c in row_identifier_cols)

        # Cek apakah kolom ini kosong total
        is_empty_col = not any(grid[r][c].strip() for r in range(header_start, max_r))
        if is_empty_col:
            continue

        if is_row_id:
            candidate_vals = []
            for r in h_rows:
                if r == col_num_row_idx:
                    continue
                v = grid[r][c].strip()
                if not v or is_year_val(v) or v.lower() in MONTH_MAP:
                    continue
                is_spanned_to_data = any(grid[r][d_col].strip() == v for d_col in range(max_c) if d_col not in row_identifier_cols)
                candidate_vals.append((v, is_spanned_to_data))
            
            specific_vals = [v for v, spanned in candidate_vals if not spanned]
            top_val = specific_vals[-1] if specific_vals else (candidate_vals[-1][0] if candidate_vals else f"Baris {col_letter}")
            ind_header, en_header = extract_cell_lang(top_val)
            sub_header = ""
            sub_header_en = ""
            sub_sub_header = ""
            sub_sub_header_en = ""
            levels = [{"ind": ind_header, "en": en_header}]
            tree_notation_ind = ind_header
            tree_notation_en = en_header
            full_title = ind_header
        else:
            levels = []
            prev_clean = None
            i = 0
            while i < len(col_h_rows):
                r = col_h_rows[i]
                val = grid[r][c].strip()
                if not val or is_bracket_col_num(val) or r == col_num_row_idx:
                    i += 1
                    continue

                ind_part, en_part = extract_cell_lang(val)

                # Jika baris berikutnya adalah baris terjemahan bahasa inggris murni
                if i + 1 < len(col_h_rows) and col_h_rows[i+1] in en_translation_rows:
                    en_val = grid[col_h_rows[i+1]][c].strip()
                    if en_val:
                        en_part = en_val
                    i += 1  # Lewati baris terjemahan

                clean_ind = ind_part.lower()
                # Hindari duplikasi berulang secara vertikal akibat merged cell / rowspan
                if clean_ind != prev_clean:
                    prev_clean = clean_ind
                    levels.append({
                        "ind": ind_part,
                        "en": en_part if en_part else ind_part
                    })
                i += 1

            if not levels:
                levels = [{"ind": f"Kolom {col_letter}", "en": f"Column {col_letter}"}]

            ind_header = levels[0]["ind"]
            en_header  = levels[0]["en"]
            sub_header = levels[1]["ind"] if len(levels) >= 2 else ""
            sub_header_en = levels[1]["en"] if len(levels) >= 2 else ""
            sub_sub_header = levels[2]["ind"] if len(levels) >= 3 else ""
            sub_sub_header_en = levels[2]["en"] if len(levels) >= 3 else ""

            tree_notation_ind = " || ".join([l["ind"] for l in levels])
            tree_notation_en  = " || ".join([l["en"] for l in levels])
            full_title = " > ".join([l["ind"] for l in levels])

        col_vals = [grid[r][c].strip() for r in range(data_start, max_r) if grid[r][c].strip()]
        preview_vals = [v for v in col_vals if v != ind_header and not is_col_num_series([v])][:6]

        columns_info.append({
            "index": c,
            "letter": col_letter,
            "header": ind_header,
            "header_en": en_header,
            "sub_header": sub_header,
            "sub_header_en": sub_header_en,
            "sub_sub_header": sub_sub_header,
            "sub_sub_header_en": sub_sub_header_en,
            "levels": levels,
            "tree_notation_ind": tree_notation_ind,
            "tree_notation_en": tree_notation_en,
            "full_title": full_title,
            "preview": preview_vals,
            "sample_count": len(col_vals),
            "is_row_identifier": is_row_id
        })

        top_baris_vals = [grid[r][best_row_col_idx].strip() for r in h_rows if grid[r][best_row_col_idx].strip()]
        baris_main_text = top_baris_vals[0].split('\n')[0].split(' / ')[0].strip().lower() if top_baris_vals else ""

        if (not is_row_id and 
            ind_header not in seen_parents and 
            not re.match(r'^Kolom\s+[A-Z0-9]+$', ind_header) and
            ind_header.lower() not in ROW_ID_KEYWORDS and
            ind_header.lower() != baris_main_text):
            seen_parents.add(ind_header)
            distinct_parents.append({
                "ind": ind_header,
                "en": en_header
            })

    # Susun matriks header grid untuk live preview
    active_indices = [c['index'] for c in columns_info]
    header_grid = []
    for r in h_rows:
        header_grid.append([grid[r][c_idx] for c_idx in active_indices])

    sample_rows = []
    sample_limit = min(data_start + 10, max_r)
    for r in range(data_start, sample_limit):
        row_cells = [grid[r][c_idx] for c_idx in active_indices]
        if any(v != "" for v in row_cells):
            sample_rows.append(row_cells)

    # 6. Susun auto_periode
    auto_periode = []
    seen_per = set()
    for c in data_cols:
        y_val = grid[year_row_idx][c].strip() if year_row_idx is not None else ""
        m_val = grid[month_row_idx][c].split('\n')[0].strip() if month_row_idx is not None else ""

        # Pastikan y_val benar-benar tahun atau periode waktu (bukan teks judul kolom)
        if y_val and not (is_year_val(y_val) or any(p in y_val.lower() for p in ['triwulan', 'tw', 'quarter', 'q'])):
            y_val = ""
        if m_val and m_val.lower() not in MONTH_MAP and not any(p in m_val.lower() for p in ['triwulan', 'tw', 'quarter', 'q']):
            m_val = ""

        per_str = ""
        per_en = ""
        if m_val and y_val:
            m_en = MONTH_MAP.get(m_val.lower(), m_val)
            per_str = f"{m_val} {y_val}"
            per_en = f"{m_en} {y_val}"
        elif y_val:
            per_str = y_val
            per_en = y_val
        elif m_val:
            m_en = MONTH_MAP.get(m_val.lower(), m_val)
            per_str = m_val
            per_en = m_en

        if per_str and per_str not in seen_per:
            seen_per.add(per_str)
            auto_periode.append({"ind": per_str, "en": per_en if per_en else per_str})

    # Fallback auto_periode dari teks judul
    if not auto_periode:
        title_texts = []
        for r in range(min(header_start, max_r)):
            title_texts.append(" ".join(grid[r][c] for c in range(min(10, max_c)) if grid[r][c]))
        sheet_clean_for_year = res.get("current_sheet") or ""
        full_title_str = " ".join(title_texts) + " " + sheet_clean_for_year
        years_found = re.findall(r'\b(20\d{2}(?:\s*[\-\/]\s*20\d{2})?)\b', full_title_str)
        if years_found:
            for y in sorted(list(set(years_found))):
                auto_periode.append({"ind": y, "en": y})

    # 7. Susun auto_rows (Mendukung hierarki majemuk, kolom kode pendamping, & multi-panel)
    auto_rows = build_dimension_rows(
        grid, best_row_col_idx, h_rows, col_num_row_idx, data_start, max_r, max_c,
        shadow_cols=shadow_cols, clean_headers=True, auto_translate_month=True,
        auto_hierarchy=True, deduplicate=True
    )

    # 8. Usulan Nama Cerdas
    best_col_meta = next((c for c in columns_info if c['index'] == best_row_col_idx), None)
    sugg_baris_ind = best_col_meta['header'] if best_col_meta else 'Judul Baris'
    sugg_baris_en  = best_col_meta['header_en'] if best_col_meta else sugg_baris_ind

    if sugg_baris_ind.lower() in ['kabupaten', 'kota', 'kab/kota']:
        sugg_baris_ind = 'Kabupaten/Kota'
        sugg_baris_en  = 'Regency/Municipality'
    elif any(k in sugg_baris_ind.lower() for k in ['tampilkan data', 'kolom', 'baris', 'lanjutan', 'no.', 'no']):
        if auto_rows:
            first_row_lower = auto_rows[0]['ind'].lower()
            if any(k in first_row_lower for k in ['cilacap', 'banyumas', 'semarang', 'kab.', 'kota']):
                sugg_baris_ind = 'Kabupaten/Kota'
                sugg_baris_en  = 'Regency/Municipality'
            elif any(k in first_row_lower for k in ['januari', 'februari', 'pebruari']):
                sugg_baris_ind = 'Bulan'
                sugg_baris_en  = 'Month'
            elif any(k in first_row_lower for k in ['imigrasi']):
                sugg_baris_ind = 'Kantor Imigrasi'
                sugg_baris_en  = 'Immigration Office'

    # Multi-Column Row Header: Jika ada kolom companion atau auto_rows mengandung ||
    has_multi_col_row = any(" || " in r.get("ind", "") for r in auto_rows)
    companion_cols = [c for c in range(best_row_col_idx + 1, first_stat_col) if c not in shadow_cols]
    comp_col = companion_cols[0] if companion_cols else None
    if (has_multi_col_row or comp_col is not None) and comp_col is not None:
        comp_meta = next((c for c in columns_info if c['index'] == comp_col), None)
        if comp_meta and comp_meta['header'] and comp_meta['header'].strip().lower() != sugg_baris_ind.strip().lower():
            comp_h_ind = comp_meta['header'].strip()
            comp_h_en  = (comp_meta['header_en'] or comp_h_ind).strip()
            if " || " not in sugg_baris_ind:
                sugg_baris_ind = f"{sugg_baris_ind} || {comp_h_ind}"
                sugg_baris_en  = f"{sugg_baris_en} || {comp_h_en}"

    title_full = " ".join(" ".join(grid[r][:10]) for r in range(min(header_start, max_r)))
    sugg_kolom_ind = ""
    sugg_kolom_en = ""

    m_menurut = re.search(r'menurut\s+([a-zA-Z\s]+?)(?:\s+dan\s+([a-zA-Z\s]+?))?(?:\s+di|\s+pada|\s+tahun|\s*,\s*20|$)', title_full, re.I)
    m_berdasarkan = re.search(r'berdasarkan\s+([a-zA-Z\s]+?)(?:\s+di|\s+pada|\s+periode|\s+tahun|\s*,\s*20|$)', title_full, re.I)
    m_per = re.search(r'per\s+([a-zA-Z\s]+?)(?:\s+di|\s+pada|\s+provinsi|\s*,\s*20|$)', title_full, re.I)

    if m_menurut:
        k1 = m_menurut.group(1).strip().title()
        k2 = m_menurut.group(2).strip().title() if m_menurut.group(2) else ""
        if any(w in k1.lower() for w in ['kabupaten', 'kecamatan', 'wilayah', 'daerah', 'bulan']) and k2:
            sugg_kolom_ind = k2
        else:
            sugg_kolom_ind = k1
    elif m_berdasarkan:
        sugg_kolom_ind = m_berdasarkan.group(1).strip().title()
    elif m_per:
        sugg_kolom_ind = m_per.group(1).strip().title()

    sheet_clean = (res.get("current_sheet") or "").strip()
    if (not sugg_kolom_ind or sugg_kolom_ind.lower() == sugg_baris_ind.lower()) and distinct_parents:
        if len(distinct_parents) == 1 and distinct_parents[0]['ind'] not in ['Kolom', 'Data']:
            sugg_kolom_ind = distinct_parents[0]['ind']
            sugg_kolom_en  = distinct_parents[0]['en']
        elif sheet_clean and sheet_clean.lower() != 'sheet1' and not re.match(r'^\d+(\.\d+)*(\s*\(\d+\))?$', sheet_clean) and not sheet_clean.lower().startswith('tabel'):
            sugg_kolom_ind = sheet_clean
            sugg_kolom_en = sheet_clean
        elif distinct_parents:
            sugg_kolom_ind = distinct_parents[0]['ind']
            sugg_kolom_en  = distinct_parents[0].get('en', sugg_kolom_ind)

    if not sugg_kolom_en:
        sugg_kolom_en = sugg_kolom_ind

    # Usulan nama kelompok periode
    sugg_periode_ind = ""
    sugg_periode_en = ""
    if auto_periode:
        if len(auto_periode) == 1:
            sugg_periode_ind = f"Tahun {auto_periode[0]['ind']}" if auto_periode[0]['ind'].isdigit() else auto_periode[0]['ind']
            sugg_periode_en = f"Year {auto_periode[0]['en']}" if auto_periode[0]['en'].isdigit() else auto_periode[0]['en']
        elif len(auto_periode) > 1:
            all_years = all(re.match(r'^(19|20)\d{2}(\s*[\-\/]\s*20\d{2})?$', p['ind']) for p in auto_periode)
            if all_years:
                sugg_periode_ind = f"Tahun {auto_periode[0]['ind']}-{auto_periode[-1]['ind']}"
                sugg_periode_en = f"Year {auto_periode[0]['en']}-{auto_periode[-1]['en']}"
            else:
                sugg_periode_ind = f"Periode {auto_periode[0]['ind']} - {auto_periode[-1]['ind']}"
                sugg_periode_en = f"Period {auto_periode[0]['en']} - {auto_periode[-1]['en']}"

    return {
        "status": "success",
        "type": res.get("type"),
        "sheets": res.get("sheets", []),
        "current_sheet": res.get("current_sheet"),
        "columns": columns_info,
        "distinct_parents": distinct_parents,
        "header_grid": header_grid,
        "sample_rows": sample_rows,
        "header_row_count": len(h_rows),
        "total_rows": max_r,
        "auto_rows": auto_rows,
        "auto_periode": auto_periode,
        "best_row_col_index": best_row_col_idx,
        "suggested_baris_ind": sugg_baris_ind,
        "suggested_baris_en": sugg_baris_en,
        "suggested_kolom_ind": sugg_kolom_ind,
        "suggested_kolom_en": sugg_kolom_en,
        "suggested_periode_ind": sugg_periode_ind,
        "suggested_periode_en": sugg_periode_en,
        "has_periode": len(auto_periode) > 0
    }

def extract_rows(source, sheet_name=None, ind_col=0, en_col=None, deduplicate=True, clean_headers=True, auto_translate_month=True, auto_hierarchy=True):
    res = fetch_raw_data(source, sheet_name)
    if "error" in res:
        return res

    grid = res.get("rows", [])
    if not grid:
        return {"error": "Sheet kosong."}

    # Unifikasi otomatis tabel multi-block (tabel berdampingan horizontal beda periode/tahun)
    grid = unify_side_by_side_blocks(grid)

    max_r = len(grid)
    max_c = max(len(r) for r in grid) if grid else 0
    header_start, header_end, data_start, h_rows, col_num_row_idx = detect_structure(grid)

    extracted = build_dimension_rows(
        grid, ind_col, h_rows, col_num_row_idx, data_start, max_r, max_c,
        en_col=en_col, clean_headers=clean_headers, auto_translate_month=auto_translate_month,
        auto_hierarchy=auto_hierarchy, deduplicate=deduplicate
    )

    return {
        "status": "success",
        "items": extracted,
        "total_extracted": len(extracted)
    }

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(json.dumps({"error": "Argumen tidak mencukupi"}, ensure_ascii=True))
        sys.exit(1)

    mode = sys.argv[1]
    src = sys.argv[2]
    sheet = sys.argv[3] if len(sys.argv) > 3 and sys.argv[3] not in ("None", "") else None

    if mode == "inspect":
        row_col_override = int(sys.argv[4]) if len(sys.argv) > 4 and sys.argv[4] not in ("None", "", "-1") else None
        result = inspect_source(src, sheet, row_col_override)
        print(json.dumps(result, ensure_ascii=True))
    elif mode == "extract":
        ind_col = int(sys.argv[4]) if len(sys.argv) > 4 else 0
        en_col = int(sys.argv[5]) if len(sys.argv) > 5 and sys.argv[5] != "None" and sys.argv[5] != "-1" else None
        dedup = (sys.argv[6].lower() in ['1', 'true', 'yes']) if len(sys.argv) > 6 else True
        clean_headers = (sys.argv[7].lower() in ['1', 'true', 'yes']) if len(sys.argv) > 7 else True
        auto_translate = (sys.argv[8].lower() in ['1', 'true', 'yes']) if len(sys.argv) > 8 else True
        auto_hierarchy = (sys.argv[9].lower() in ['1', 'true', 'yes']) if len(sys.argv) > 9 else True
        result = extract_rows(src, sheet, ind_col, en_col, deduplicate=dedup, clean_headers=clean_headers, auto_translate_month=auto_translate, auto_hierarchy=auto_hierarchy)
        print(json.dumps(result, ensure_ascii=True))
    else:
        print(json.dumps({"error": f"Mode tidak dikenal: {mode}"}, ensure_ascii=True))

