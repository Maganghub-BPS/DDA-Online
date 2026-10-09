#Untuk mengunduh dan melihat spreadsheet, Data API 

import re
import csv
import ssl
import time
import json
import urllib.request
import urllib.error
from io import StringIO
from config import SATUDATA_TOKEN, HTTP_TIMEOUT, HTTP_USER_AGENT
from normalizer import parse_num

def expand_merged_cells_row(row):
    if not isinstance(row, list): return row
    new_row = []
    for cell in row:
        c_str = str(cell).strip()
        if re.search(r'\s{2,}|\t', c_str):
            parts = [p.strip() for p in re.split(r'\s{2,}|\t', c_str) if p.strip()]
            if len(parts) >= 2 and (
                all(parse_num(p) is not None for p in parts) or
                all(re.match(r'^(?:19|20)\d{2}(?:/\d{4})?$', p) for p in parts) or
                all(re.match(r'^\(?\d{1,2}\)?$', p) for p in parts)
            ):
                new_row.extend(parts)
                continue
        new_row.append(cell)
    return new_row

def fetch_data_from_link(link: str):
    if not link or not isinstance(link, str):
        return {"error": "Tautan data kosong", "is_link_empty": True}
    link = link.strip('\r\n\t')
    if not link:
        return {"error": "Tautan data kosong", "is_link_empty": True}

    ssl_ctx = ssl._create_unverified_context()

    try:
        # 0. Local Excel Files
        if '.xlsx' in link or '.xls' in link:
            import openpyxl
            try:
                sheet_name = None
                file_path = link
                if '|' in link:
                    file_path, sheet_name = link.split('|', 1)
                
                wb = openpyxl.load_workbook(file_path, data_only=True)
                
                if sheet_name and sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                else:
                    sheet = wb.active
                    for sn in wb.sheetnames:
                        if 'Tabel' in sn or 'Table' in sn:
                            sheet = wb[sn]
                            break
                
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(expand_merged_cells_row([str(c).strip() if c is not None else "" for c in row]))
                return rows
            except Exception as e:
                return {"error": f"Gagal membaca file Excel lokal: {str(e)}"}

        # 1. Google Spreadsheets
        if "docs.google.com/spreadsheets" in link:
            match_key = re.search(r"/d/([a-zA-Z0-9-_]+)", link)
            doc_key = match_key.group(1) if match_key else None
            gid = None
            if "gid=" in link:
                gid = link.split("gid=")[1].split("#")[0].split("&")[0].strip()

            now_ts = int(time.time())
            if doc_key:
                urls_to_try = [
                    f"https://docs.google.com/spreadsheets/d/{doc_key}/export?format=csv" + (f"&gid={gid}" if gid else "") + f"&_t={now_ts}",
                    f"https://docs.google.com/spreadsheets/d/{doc_key}/gviz/tq?tqx=out:csv" + (f"&gid={gid}" if gid else "") + f"&_t={now_ts}"
                ]
            else:
                csv_link = re.sub(r"/edit.*", "/export?format=csv", link)
                if gid: csv_link += f"&gid={gid}"
                urls_to_try = [csv_link + f"&_t={now_ts}"]
            
            max_retries = 3
            last_err = None
            for csv_url in urls_to_try:
                for attempt in range(max_retries):
                    try:
                        req = urllib.request.Request(
                            csv_url,
                            headers={"User-Agent": HTTP_USER_AGENT, "Cache-Control": "no-cache", "Pragma": "no-cache"}
                        )
                        with urllib.request.urlopen(req, timeout=HTTP_TIMEOUT, context=ssl_ctx) as response:
                            csv_data = response.read().decode("utf-8", errors="ignore")
                            if csv_data.strip().startswith(("<html", "<!DOCTYPE", "<!doctype", "<head")):
                                last_err = "Mendapat respon HTML/Sinkhole dari jaringan, bukan CSV"
                                break
                            reader = csv.reader(StringIO(csv_data))
                            return [expand_merged_cells_row(row) for row in reader]
                    except Exception as e:
                        last_err = e
                        if attempt < max_retries - 1:
                            time.sleep(0.5 * (attempt + 1))
            return {"error": f"Gagal mengambil data dari Google Sheets: {str(last_err)}"}
                    
        # 2. API Portal Satu Data
        elif "view_portal_tabel" in link:
            if "?id=" in link:
                ids = link.split("?id=")[1]
                id_p = ids.split(",")[0].strip()
                api_url = f"https://satudata.jatengprov.go.id/api/v1/data/{id_p}?peruntukan=DDA"
                token = SATUDATA_TOKEN
                
                req = urllib.request.Request(api_url, headers={
                    "User-Agent": HTTP_USER_AGENT, "Authorization": f"Bearer {token}", "Accept": "application/json"
                })
                try:
                    with urllib.request.urlopen(req, timeout=HTTP_TIMEOUT, context=ssl_ctx) as response:
                        data = json.loads(response.read().decode("utf-8"))
                        records = data.get("data", data)
                        if not records or len(records) == 0:
                            return {"error": "Tidak ada data dalam API", "is_api_empty": True}
                        return records
                except urllib.error.HTTPError as e:
                    return {"error": "Tidak ada data dalam API", "is_api_empty": True, "detail": f"HTTP Error {e.code}: {e.reason}"}
                except Exception as e:
                    return {"error": "Tidak ada data dalam API", "is_api_empty": True, "detail": str(e)}
        
        return {"error": "Format link tidak didukung atau butuh login (misal: Simdasi)", "is_link_unsupported": True}
    except Exception as e:
        return {"error": str(e)}
