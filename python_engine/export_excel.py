#Untuk memformat ke excel hasil dari perbandingan

import sys
import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from compare_table import compare_head_to_head

# PENTING: Menerima parameter id_db
def export_table_to_excel(pdf_path, db_json_path, target_table, output_path, tolerance=0.0, id_db=None):
    res = compare_head_to_head(pdf_path, target_table, db_json_path, tolerance=tolerance, id_db=id_db)
    if res.get("status") != "success":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Tabel {target_table}"
        ws['A1'] = f"Status Validasi Tabel {target_table}: {res.get('message', 'Error')}"
        wb.save(output_path)
        return output_path
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Tabel {target_table}"
    
    title_font = Font(name="Calibri", size=14, bold=True, color="1F2937")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    
    diff_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    match_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    pdf_only_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'), right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'), bottom=Side(style='thin', color='E5E7EB')
    )
    
    ws['A1'] = f"BERITA ACARA REKONSILIASI DATA - TABEL {target_table}"
    ws['A1'].font = title_font
    
    summary = res.get("summary", {})
    total_diffs = summary.get("total_diffs", len(res.get("diffs", [])))
    total_matches = summary.get("total_matches", len(res.get("matches", [])))
    total_pdf_only = summary.get("total_pdf_only", len(res.get("pdf_only", [])))
    
    sub_txt = f"Sumber: {res.get('source_type', 'Database Sumber')} | Total Cocok: {total_matches} | Selisih: {total_diffs} | Khusus PDF: {total_pdf_only}"
    ws['A2'] = sub_txt
    ws['A2'].font = Font(name="Calibri", size=10, italic=True, color="4B5563")
    
    headers = ["No", "Wilayah / Baris", "Bagian / Kategori", "Nama Metrik / Data", "Nilai di Buku PDF", "Nilai di Database", "Selisih", "Status Validasi", "Keterangan Catatan"]
    row_idx = 4
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row_idx].height = 24
    
    row_idx += 1
    counter = 1
    
    for d in res.get("diffs", []):
        ws.cell(row=row_idx, column=1, value=counter).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=d.get("wilayah", ""))
        ws.cell(row=row_idx, column=3, value=d.get("section", "-"))
        ws.cell(row=row_idx, column=4, value=d.get("metric", ""))
        ws.cell(row=row_idx, column=5, value=d.get("pdf_val", "")).alignment = Alignment(horizontal="right")
        ws.cell(row=row_idx, column=6, value=d.get("src_val", "")).alignment = Alignment(horizontal="right")
        ws.cell(row=row_idx, column=7, value=d.get("diff", "")).alignment = Alignment(horizontal="right")
        
        status_cell = ws.cell(row=row_idx, column=8, value="BEDA NILAI")
        status_cell.font = Font(name="Calibri", bold=True, color="991B1B")
        status_cell.alignment = Alignment(horizontal="center")
        
        ws.cell(row=row_idx, column=9, value=d.get("note", ""))
        for c in range(1, 10):
            ws.cell(row=row_idx, column=c).fill = diff_fill
            ws.cell(row=row_idx, column=c).border = thin_border
        row_idx += 1
        counter += 1
        
    for p in res.get("pdf_only", []):
        ws.cell(row=row_idx, column=1, value=counter).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=p.get("wilayah", ""))
        ws.cell(row=row_idx, column=3, value=p.get("section", "-"))
        ws.cell(row=row_idx, column=4, value=p.get("metric", ""))
        val_str = ' | '.join(str(x) for x in p.get("values", []))
        ws.cell(row=row_idx, column=5, value=val_str).alignment = Alignment(horizontal="right")
        ws.cell(row=row_idx, column=6, value="-").alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=7, value="-").alignment = Alignment(horizontal="center")
        
        status_cell = ws.cell(row=row_idx, column=8, value="HANYA DI PDF")
        status_cell.font = Font(name="Calibri", bold=True, color="92400E")
        status_cell.alignment = Alignment(horizontal="center")
        
        ws.cell(row=row_idx, column=9, value=p.get("note", ""))
        for c in range(1, 10):
            ws.cell(row=row_idx, column=c).fill = pdf_only_fill
            ws.cell(row=row_idx, column=c).border = thin_border
        row_idx += 1
        counter += 1
        
    tol_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    for m in res.get("matches", []):
        ws.cell(row=row_idx, column=1, value=counter).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=m.get("wilayah", ""))
        ws.cell(row=row_idx, column=3, value=m.get("section", "-"))
        ws.cell(row=row_idx, column=4, value=m.get("metric", ""))
        ws.cell(row=row_idx, column=5, value=m.get("pdf_val", m.get("val", ""))).alignment = Alignment(horizontal="right")
        ws.cell(row=row_idx, column=6, value=m.get("src_val", m.get("val", ""))).alignment = Alignment(horizontal="right")
        
        is_tol = m.get("is_tolerance_match", False)
        tol_diff = m.get("tolerance_diff", 0)
        ws.cell(row=row_idx, column=7, value=tol_diff if is_tol else 0).alignment = Alignment(horizontal="right")
        
        status_label = "COCOK (TOLERANSI)" if is_tol else "COCOK"
        status_cell = ws.cell(row=row_idx, column=8, value=status_label)
        status_cell.font = Font(name="Calibri", bold=True, color="0369A1" if is_tol else "166534")
        status_cell.alignment = Alignment(horizontal="center")
        
        note_txt = f"Sesuai toleransi ({'+' if tol_diff > 0 else ''}{tol_diff})" if is_tol else "Sesuai persis"
        ws.cell(row=row_idx, column=9, value=note_txt)
        
        current_fill = tol_fill if is_tol else match_fill
        for c in range(1, 10):
            ws.cell(row=row_idx, column=c).fill = current_fill
            ws.cell(row=row_idx, column=c).border = thin_border
        row_idx += 1
        counter += 1
        
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 10)
        
    wb.save(output_path)
    return output_path

if __name__ == "__main__":
    if len(sys.argv) < 5:
        print("Usage: python export_excel.py <pdf_path> <db_json_path> <target_table> <output_path> [tolerance] [id_db]")
        sys.exit(1)
        
    pdf_p = sys.argv[1]
    db_p = sys.argv[2]
    tbl = sys.argv[3]
    out_p = sys.argv[4]
    tol = float(sys.argv[5]) if len(sys.argv) > 5 else 0.0
    id_db = sys.argv[6] if len(sys.argv) > 6 else None
    
    export_table_to_excel(pdf_p, db_p, tbl, out_p, tolerance=tol, id_db=id_db)
    print("Export complete")